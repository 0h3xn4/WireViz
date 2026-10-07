"""CSV/XLSX import of interfaces. Plans first: nothing changes until the user confirms.

`plan_interface_import` validates every row against the project (and against the rows before it)
and returns per-row results plus the operations for the good rows, applied as one undo step.
"""

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path

from . import edit
from .commands import Op
from .errors import HarnessError
from .ids import ID_RE
from .model import Project

MAX_ROWS = 20000
MAX_BYTES = 8 * 1024 * 1024
FIELDS: tuple[tuple[str, str], ...] = (
    ("id", "Interface ID"),
    ("type", "Interface type"),
    ("from", "From unit"),
    ("to", "To unit"),
    ("redundancy", "Redundancy"),
)
_ALIASES = {
    "id": ("interface", "id", "interfaceid", "ifid", "name"),
    "type": ("type", "interfacetype", "kind", "protocol"),
    "from": ("from", "fromunit", "source", "src", "a"),
    "to": ("to", "tounit", "destination", "dest", "target", "b"),
    "redundancy": ("redundancy", "chain", "side", "redundant"),
}


class ImportError_(HarnessError):  # noqa: N801 (avoid shadowing the builtin ImportError)
    """The file cannot be read as a table."""


Table = list[list[str]]


def parse_csv(text: str) -> Table:
    if len(text) > MAX_BYTES:
        raise ImportError_("The file is too large to import (limit 8 MB).")
    rows = [
        [c.strip() for c in r] for r in csv.reader(io.StringIO(text)) if any(c.strip() for c in r)
    ]
    if len(rows) > MAX_ROWS + 1:
        raise ImportError_(f"The file has more than {MAX_ROWS} rows.")
    return rows


def read_table(path: Path | str) -> Table:
    """Read a .csv or .xlsx file into rows of text. The first row is the header."""
    path = Path(path)
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise ImportError_(f"The file could not be read ({exc.strerror}).") from exc
    if size > MAX_BYTES:
        raise ImportError_("The file is too large to import (limit 8 MB).")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        try:
            return parse_csv(path.read_bytes().decode("utf-8-sig"))
        except UnicodeDecodeError as exc:
            raise ImportError_(
                "The file is not UTF-8 text. Save it as CSV UTF-8 and try again."
            ) from exc
        except OSError as exc:
            raise ImportError_(f"The file could not be read ({exc.strerror}).") from exc
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    raise ImportError_("Only .csv and .xlsx files can be imported.")


def _read_xlsx(path: Path) -> Table:
    import zipfile

    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except (OSError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ImportError_("The file is not a valid .xlsx workbook.") from exc
    try:
        ws = wb.active
        if ws is None:
            raise ImportError_("The workbook has no sheet.")
        rows: Table = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c).strip() for c in row]
            if any(cells):
                rows.append(cells)
            if len(rows) > MAX_ROWS + 1:
                raise ImportError_(f"The sheet has more than {MAX_ROWS} rows.")
        return rows
    finally:
        wb.close()


def guess_mapping(header: list[str]) -> dict[str, int]:
    """Best column for each field, by header name; unknown fields default to the field's position."""
    norm = ["".join(ch for ch in h.lower() if ch.isalnum()) for h in header]
    result: dict[str, int] = {}
    used: set[int] = set()
    for key, _ in FIELDS:
        for alias in _ALIASES[key]:
            hit = next((i for i, h in enumerate(norm) if h == alias and i not in used), None)
            if hit is not None:
                result[key] = hit
                used.add(hit)
                break
    for pos, (key, _) in enumerate(FIELDS):
        if key not in result and pos < len(header) and pos not in used:
            result[key] = pos
            used.add(pos)
    return result


@dataclass(frozen=True)
class RowResult:
    row_number: int  # 1-based, counting the header as row 1
    ok: bool
    message: str
    values: tuple[str, ...]


@dataclass
class ImportPlan:
    rows: list[RowResult] = field(default_factory=list)
    ops: list[Op] = field(default_factory=list)

    @property
    def ok_count(self) -> int:
        return sum(r.ok for r in self.rows)

    @property
    def error_count(self) -> int:
        return len(self.rows) - self.ok_count


def _cell(row: list[str], mapping: dict[str, int], key: str) -> str:
    idx = mapping.get(key)
    return row[idx].strip() if idx is not None and idx < len(row) else ""


def _match_type(project: Project, text: str) -> str | None:
    low = text.casefold()
    for tid, t in sorted(project.interface_types.items()):
        if tid.casefold() == low or t.name.casefold() == low:
            return tid
    return None


def _match_unit(project: Project, text: str) -> str | None:
    low = text.casefold()
    for uid, u in sorted(project.units.items()):
        if uid.casefold() == low or u.name.casefold() == low:
            return uid
    return None


def plan_interface_import(project: Project, table: Table, mapping: dict[str, int]) -> ImportPlan:
    """Validate each data row; good rows become operations. The project itself is not changed."""
    plan = ImportPlan()
    if not table:
        return plan
    scratch = project
    seen = {i.casefold() for i in project.interfaces}
    for n, row in enumerate(table[1:], start=2):
        vals = tuple(_cell(row, mapping, k) for k, _ in FIELDS)
        iid, ty, a, b, red = vals
        error = ""
        tid = _match_type(project, ty)
        ua, ub = _match_unit(project, a), _match_unit(project, b)
        if not ID_RE.fullmatch(iid) or iid.endswith("."):
            error = f"Invalid interface ID '{iid}'"
        elif iid.casefold() in seen:
            error = f"ID {iid} is already used"
        elif tid is None:
            error = f"Unknown interface type '{ty}'"
        elif ua is None:
            error = f"Unit '{a}' does not exist"
        elif ub is None:
            error = f"Unit '{b}' does not exist"
        elif ua == ub:
            error = "A unit cannot be connected to itself"
        elif red and red.casefold() not in ("nominal", "redundant", "none"):
            error = f"Redundancy must be nominal, redundant or none, not '{red}'"
        if not error and tid is not None and ua is not None and ub is not None:
            try:
                ops, _ = edit.ops_add_interface(
                    scratch, tid, ua, ub, interface_id=iid, redundancy=red.casefold() or None
                )
            except edit.EditError as exc:
                error = str(exc)
            else:
                scratch = edit.clone_with(scratch, ops)
                plan.ops.extend(ops)
                seen.add(iid.casefold())
        plan.rows.append(RowResult(n, not error, error, vals))
    return plan
